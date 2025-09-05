# chatbot_factory/routes/knowledge_routes.py
from flask import render_template, url_for, flash, redirect, request, Blueprint, abort, jsonify
from flask_login import current_user, login_required
from flask_babel import gettext as _
from .. import db
from ..models import Bot, KnowledgeBase, Product
from ..forms import KnowledgeBaseForm, ProductForm
from ..services.ai_service import ai_service
import openpyxl
import os
import logging
from werkzeug.utils import secure_filename
import docx

knowledge_bp = Blueprint('knowledge', __name__)

@knowledge_bp.route("/bot/<int:bot_id>/knowledge", methods=['GET'])
@login_required
def manage_knowledge(bot_id):
    bot = Bot.query.get_or_404(bot_id)
    if bot.owner != current_user:
        abort(403)
    
    text_form = KnowledgeBaseForm()
    entries = bot.knowledge_base
    products = bot.products
    return render_template('knowledge_base.html', title=_('Knowledge Base'), bot=bot, 
                           text_form=text_form, entries=entries, products=products)

@knowledge_bp.route("/bot/<int:bot_id>/chat", methods=['POST'])
@login_required
def bot_chat(bot_id):
    """Chat endpoint for testing bot responses"""
    bot = Bot.query.get_or_404(bot_id)
    if bot.owner != current_user:
        abort(403)
    
    data = request.get_json()
    if not data or 'message' not in data:
        return jsonify({'error': 'Message is required'}), 400
    
    user_message = data['message']
    
    # Build system prompt with knowledge base
    knowledge_context = ""
    
    # Add text entries
    if bot.knowledge_base:
        knowledge_context += "\n\n--- UMUMIY MA'LUMOTLAR BAZASI ---\n"
        for entry in bot.knowledge_base:
            knowledge_context += f"MAVZU: {entry.title}\nMA'LUMOT: {entry.content}\n\n"
    
    # Add products
    if bot.products:
        knowledge_context += "\n\n--- MAHSULOTLAR RO'YXATI ---\n"
        for product in bot.products:
            knowledge_context += f"MAHSULOT NOMI: {product.name}\n"
            if product.price:
                knowledge_context += f"NARXI: {product.price}\n"
            if product.description:
                knowledge_context += f"TAVSIFI: {product.description}\n"
            if product.image_url:
                knowledge_context += f"RASM HAVOLASI: {product.image_url}\n"
            knowledge_context += "---\n"

    if knowledge_context:
        knowledge_context += "--- BILIMLAR BAZASI TUGADI ---\n"
        knowledge_context += "Foydalanuvchi savoliga javob berish uchun YUQORIDAGI BILIMLAR BAZASIDAN (Umumiy ma'lumotlar va Mahsulotlar) birinchi navbatda foydalaning. Agar savol bu ma'lumotlarga aloqador bo'lmasa, umumiy bilimingizdan foydalanib javob bering."

    system_prompt = f"{bot.system_prompt}\nSizning ismingiz \"{bot.name}\".\n{knowledge_context}"
    
    # Get AI response
    response_data = ai_service.generate_response(
        prompt=user_message,
        system_prompt=system_prompt,
        bot=bot
    )
    
    return jsonify(response_data)

@knowledge_bp.route("/bot/<int:bot_id>/knowledge/add_text", methods=['POST'])
@login_required
def add_text_entry(bot_id):
    bot = Bot.query.get_or_404(bot_id)
    if bot.owner != current_user:
        abort(403)
    form = KnowledgeBaseForm()
    if form.validate_on_submit():
        entry = KnowledgeBase(title=form.title.data, content=form.content.data, bot=bot)
        db.session.add(entry)
        db.session.commit()
        flash(_('New text entry has been added!'), 'success')
    return redirect(url_for('knowledge.manage_knowledge', bot_id=bot_id))

@knowledge_bp.route("/bot/<int:bot_id>/add_product", methods=['GET', 'POST'])
@login_required
def add_product(bot_id):
    bot = Bot.query.get_or_404(bot_id)
    if bot.owner != current_user:
        abort(403)
    form = ProductForm()
    if form.validate_on_submit():
        product = Product(bot_id=bot.id, name=form.name.data, price=form.price.data, 
                          description=form.description.data, image_url=form.image_url.data)
        db.session.add(product)
        db.session.commit()
        flash(_('New product has been added to the knowledge base!'), 'success')
        return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))
    return render_template('add_product.html', title=_('Add Product'), form=form, bot=bot)

@knowledge_bp.route("/knowledge/<int:entry_id>/delete", methods=['POST'])
@login_required
def delete_knowledge(entry_id):
    # Bu funksiya o'zgarishsiz qoladi
    entry = KnowledgeBase.query.get_or_404(entry_id)
    bot = entry.bot
    if bot.owner != current_user:
        abort(403)
    db.session.delete(entry)
    db.session.commit()
    flash(_('The text entry has been deleted.'), 'success')
    return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))

@knowledge_bp.route("/product/<int:product_id>/delete", methods=['POST'])
@login_required
def delete_product(product_id):
    product = Product.query.get_or_404(product_id)
    bot = product.bot_owner
    if bot.owner != current_user:
        abort(403)
    db.session.delete(product)
    db.session.commit()
    flash(_('The product has been deleted.'), 'success')
    return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))

@knowledge_bp.route("/bot/<int:bot_id>/knowledge/upload_products", methods=['POST'])
@login_required
def upload_products(bot_id):
    bot = Bot.query.get_or_404(bot_id)
    if bot.owner != current_user:
        abort(403)

    if 'file' not in request.files:
        flash(_('No file part'), 'danger')
        return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))
    
    file = request.files['file']
    if file.filename == '':
        flash(_('No selected file'), 'danger')
        return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))

    if file and file.filename.endswith('.xlsx'):
        try:
            workbook = openpyxl.load_workbook(file)
            sheet = workbook.active
            
            product_count = 0
            # Birinchi qator (sarlavha)ni o'tkazib yuborish uchun min_row=2
            for row in sheet.iter_rows(min_row=2, values_only=True):
                name, price, description, image_url = row[0], row[1], row[2], row[3]
                if name: # Faqat nomi bor qatorlarni qo'shamiz
                    product = Product(bot_id=bot.id, name=name, price=price, 
                                      description=description, image_url=image_url)
                    db.session.add(product)
                    product_count += 1
            
            db.session.commit()
            flash(_('%(num)s products have been successfully uploaded!', num=product_count), 'success')
        except Exception as e:
            db.session.rollback()
            flash(_('An error occurred while processing the file. Please check the format. Error: %(error)s', error=str(e)), 'danger')
    else:
        flash(_('Invalid file type. Please upload a .xlsx file.'), 'danger')
        
    return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))

@knowledge_bp.route("/bot/<int:bot_id>/knowledge/upload_text_file", methods=['POST'])
@login_required
def upload_text_file(bot_id):
    bot = Bot.query.get_or_404(bot_id)
    if bot.owner != current_user:
        abort(403)

    if 'file' not in request.files:
        flash(_('No file part'), 'danger')
        return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))
    
    file = request.files['file']
    if file.filename == '':
        flash(_('No selected file'), 'danger')
        return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))

    if file:
        filename = secure_filename(file.filename)
        content = ""
        try:
            # Fayl kengaytmasini tekshirish
            if filename.endswith('.txt'):
                # UTF-8 da o'qishga harakat qilish
                content = file.read().decode('utf-8', errors='ignore')
            elif filename.endswith('.docx'):
                doc = docx.Document(file)
                full_text = [para.text for para in doc.paragraphs]
                content = '\n'.join(full_text)
            else:
                flash(_('Invalid file type. Please upload a .txt or .docx file.'), 'danger')
                return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))

            # Agar kontent mavjud bo'lsa, bazaga saqlash
            if content.strip():
                title = os.path.splitext(filename)[0]
                entry = KnowledgeBase(title=title, content=content, bot=bot)
                db.session.add(entry)
                db.session.commit()
                flash(_('File "%(name)s" has been successfully processed and added.', name=filename), 'success')
            else:
                flash(_('Could not extract any text from the file "%(name)s".', name=filename), 'warning')

        except Exception as e:
            # Foydalanuvchiga tushunarli xabar berish
            logging.error(f"File upload error for user {current_user.id}: {e}") # Logga to'liq xatoni yozish
            db.session.rollback()
            flash(_('An unexpected error occurred while processing the file. Please ensure it is a valid and properly encoded text file.'), 'danger')
        
    return redirect(url_for('knowledge.manage_knowledge', bot_id=bot.id))